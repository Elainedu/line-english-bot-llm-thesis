from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 創建一個新的Word文件
doc = Document()

# 設定中文字體
def set_chinese_font(run):
    run.font.name = '微軟正黑體'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

# 添加標題
title = doc.add_heading('表4-1d 聽力練習難度級別對照表', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    set_chinese_font(run)
    run.font.size = Pt(16)
    run.bold = True

# 準備表格內容
headers = ["難度級別", "詞彙複雜度", "句型結構", "語速/流暢度", "互動複雜度"]
content = [
    ["初級\n(Beginner)", "基礎詞彙\n10-30個單詞\n基本日常主題", "簡單直接句型\n基本表達", "緩慢清晰", "單一簡單問題\n直接信息提取\n0-100分評分"],
    ["中級\n(Intermediate)", "常用表達\n30-50個單詞\n擴展生活主題", "多樣化句型\n連接詞使用", "中等語速", "單一中等問題\n間接信息提取\n具體回饋"],
    ["高級\n(Advanced)", "自然語言表達\n50-80個單詞\n專業或抽象主題", "複雜句式結構\n多層次表達", "自然語速", "多個深入問題(3個)\n綜合理解能力\n原文提供"]
]

# 創建表格 (先創建1行表頭)
table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'

# 設置表格寬度
table.autofit = False
table.allow_autofit = False

# 添加表頭
header_cells = table.rows[0].cells
for i, header in enumerate(headers):
    header_cells[i].text = header
    for paragraph in header_cells[i].paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_chinese_font(run)
            run.bold = True

# 添加內容行
for row_content in content:
    # 添加新行
    row_cells = table.add_row().cells
    
    # 填充行內容
    for i, text in enumerate(row_content):
        row_cells[i].text = text
        # 設置字體
        for paragraph in row_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                set_chinese_font(run)
                if i == 0:  # 第一列（難度級別）加粗
                    run.bold = True

# 調整列寬
table.columns[0].width = Cm(2.5)  # 難度級別列
table.columns[1].width = Cm(4.5)  # 詞彙複雜度列
table.columns[2].width = Cm(3.5)  # 句型結構列
table.columns[3].width = Cm(2.5)  # 語速/流暢度列
table.columns[4].width = Cm(4.5)  # 互動複雜度列

# 儲存檔案
file_path = r"C:\Users\dyl89\OneDrive - 國立高雄科技大學\桌面\論文、研討會\論文\04論文書面資料\圖片\聽力練習難度級別對照表.docx"
doc.save(file_path)
print(f"文件已成功儲存至: {file_path}")